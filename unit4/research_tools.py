"""Free research tools and readers for files supplied with a question."""
import csv
import io
import ipaddress
import json
from pathlib import Path
import socket
from urllib.parse import urljoin, urlparse, parse_qs

import requests
from bs4 import BeautifulSoup
from smolagents import tool
from smolagents.default_tools import PythonInterpreterTool


def public_url(url):
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("Use a public http or https URL.")
    for record in socket.getaddrinfo(parsed.hostname, parsed.port or 443, type=socket.SOCK_STREAM):
        if not ipaddress.ip_address(record[4][0]).is_global:
            raise ValueError("Local and private network addresses are not allowed.")
    return url


def fetch_page(url):
    for _ in range(5):
        public_url(url)
        with requests.get(url, timeout=20, allow_redirects=False, stream=True,
                          headers={"User-Agent": "AlfredCourseAgent/1.0"}) as response:
            if response.status_code in {301, 302, 303, 307, 308}:
                url = urljoin(url, response.headers["Location"])
                continue
            response.raise_for_status()
            content = bytearray()
            for chunk in response.iter_content(65536):
                content.extend(chunk)
                if len(content) > 4 * 1024 * 1024:
                    raise ValueError("Page is too large; use a more specific source.")
            return url, bytes(content), response.headers.get("Content-Type", "")
    raise ValueError("Too many page redirects.")


def file_in(directory, filename):
    root = Path(directory).resolve()
    path = (root / filename).resolve()
    if path.parent != root or not path.is_file():
        raise ValueError("Choose a file supplied with this question.")
    return path


def extract_file(path):
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(f"Page {i + 1}: {p.extract_text() or ''}" for i, p in enumerate(PdfReader(path).pages))
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, data_only=True, read_only=True)
        try:
            parts = []
            for sheet in book:
                parts.append(f"Sheet: {sheet.title}")
                for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    parts.append(f"Row {i}: " + json.dumps(row, default=str))
                    if i >= 3000:
                        parts.append("Stopped at row 3000.")
                        break
            return "\n".join(parts)
        finally:
            book.close()
    if suffix in {".txt", ".py", ".csv", ".json", ".md", ".html"}:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError("This reader supports text, Python source, CSV, JSON, PDF, and XLSX files. Use transcription for audio. Image/video scene understanding is not supported.")


def build_tools(directory, api_key):
    @tool
    def web_search(query: str) -> str:
        """Search public sources for facts. Results include links to inspect.

        Args:
            query: A focused search query, without answer keys or benchmark solutions.
        """
        from ddgs import DDGS
        results = list(DDGS(timeout=15).text(query, max_results=4))
        return json.dumps([{"title": r.get("title"), "url": r.get("href"),
                            "snippet": r.get("body", "")[:650]} for r in results])

    @tool
    def read_webpage(url: str, offset: int = 0) -> str:
        """Read a URL from search results or the question. For the next part,
        repeat the same url with a larger offset. There is no browser cursor.

        Args:
            url: Required full http or https source URL, including on repeat calls.
            offset: Starting character position, default 0. Use offset, not loc or cursor.
        """
        final_url, data, kind = fetch_page(url)
        if "pdf" in kind.lower() or data.startswith(b"%PDF"):
            from pypdf import PdfReader
            text = "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
        else:
            soup = BeautifulSoup(data, "html.parser")
            for element in soup(["script", "style", "nav", "footer"]):
                element.decompose()
            text = soup.get_text("\n", strip=True)
        offset = max(0, offset)
        return json.dumps({"source": final_url, "total_characters": len(text),
                           "offset": offset, "text": text[offset:offset + 3500],
                           "next_call": {"url": final_url, "offset": offset + 3500}
                           if offset + 3500 < len(text) else None})

    @tool
    def read_attachment(filename: str, offset: int = 0) -> str:
        """Read text, Python source, PDF, or spreadsheet rows from the task file.

        Args:
            filename: Exact attachment filename provided in the question context.
            offset: Starting character position; increase to continue the file.
        """
        text = extract_file(file_in(directory, filename))
        offset = max(0, offset)
        return json.dumps({"filename": filename, "total_characters": len(text),
                           "offset": offset, "text": text[offset:offset + 5000]})

    @tool
    def transcribe_audio(filename: str) -> str:
        """Transcribe a supplied audio file through Groq's free-plan audio API.

        Args:
            filename: Exact name of the task's MP3, WAV, M4A, OGG, FLAC, or WEBM file.
        """
        from openai import OpenAI
        path = file_in(directory, filename)
        if path.suffix.lower() not in {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".webm"}:
            raise ValueError("Choose a supported audio attachment.")
        client = OpenAI(base_url="https://api.groq.com/openai/v1", api_key=api_key, max_retries=0, timeout=60)
        with path.open("rb") as stream:
            result = client.audio.transcriptions.create(model="whisper-large-v3-turbo", file=stream)
        return result.text[:10000]

    @tool
    def youtube_captions(url: str) -> str:
        """Read available YouTube captions; cannot inspect visual scenes.

        Args:
            url: YouTube video URL from the question.
        """
        from youtube_transcript_api import YouTubeTranscriptApi
        parsed = urlparse(url)
        if parsed.hostname in {"youtu.be", "www.youtu.be"}:
            video_id = parsed.path.strip("/")
        elif parsed.hostname in {"youtube.com", "www.youtube.com"}:
            video_id = parse_qs(parsed.query).get("v", [""])[0]
        else:
            raise ValueError("Use a YouTube URL.")
        if not video_id or len(video_id) > 20:
            raise ValueError("Invalid video ID.")
        captions = YouTubeTranscriptApi().fetch(video_id, languages=["en"])
        return " ".join(p.text for p in captions)[:10000]

    return [web_search, read_webpage, read_attachment, transcribe_audio, youtube_captions,
            PythonInterpreterTool(authorized_imports=["math", "statistics", "itertools", "collections"], timeout_seconds=5)]
